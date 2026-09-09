"""Operator-facing machine plan workbook.

Presentation only: every quantity comes from existing engine/scheduler results.
No planning, persistence, cache, routing, or Google Sheets mutation occurs here.
"""
from __future__ import annotations

import datetime as dt
import io
import os
from types import SimpleNamespace
from collections import defaultdict
from dataclasses import dataclass
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from mp_seed import norm_code


NAVY = "1F3864"
TERRA = "C55A11"
DAY_FILL = PatternFill("solid", fgColor="FFF6EC")
NIGHT_FILL = PatternFill("solid", fgColor="E8EDF5")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TOTAL_FILL = PatternFill("solid", fgColor=TERRA)
THIN = Side(style="thin", color="D9DEE8")
HEAVY = Side(style="medium", color=NAVY)


class OperatorExportInvariantError(ValueError):
    """The scheduler result cannot be represented in operator block format."""


@dataclass(frozen=True)
class OperatorBlock:
    machine: str
    plan_type: str
    day: int
    date: dt.date
    shift: str
    block_hours: int
    item_code: str
    raw_code: str
    item_name: str
    material: str
    target_pieces: float
    kg: float
    excess_hours: float


def _plan_dates(total_days: int, month: str) -> dict[int, dt.date]:
    year, month_number = (int(value) for value in month.split("-"))
    current = dt.date(year, month_number, 1)
    dates: list[dt.date] = []
    while len(dates) < total_days:
        if current.weekday() != 6:
            dates.append(current)
        current += dt.timedelta(days=1)
    return {index + 1: value for index, value in enumerate(dates)}


def _source_maps(source_rows: Optional[Iterable]) -> tuple[dict, dict]:
    names: dict[str, str] = {}
    raw_codes: dict[str, str] = {}
    for row in source_rows or []:
        raw = str(
            getattr(row, "raw_code", None)
            or (row.get("raw_code") if isinstance(row, dict) else "")
            or (row.get("item_code") if isinstance(row, dict) else "")
            or ""
        ).strip()
        key = norm_code(raw)
        if not key:
            continue
        name = str(
            getattr(row, "item_name", None)
            or (row.get("item_name") if isinstance(row, dict) else "")
            or ""
        ).strip()
        raw_codes.setdefault(key, raw)
        if name:
            names.setdefault(key, name)
    return names, raw_codes


def plan_run_source_rows(run_row: Optional[dict]) -> list:
    """Rebuild display-only source rows from stored demand and original workbook."""
    if not run_row:
        return []
    names: dict[str, str] = {}
    path = str(run_row.get("uploaded_file_path") or "")
    if path and os.path.isfile(path):
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            for worksheet in workbook.worksheets:
                header_row = 0
                code_col = name_col = None
                for row_number, values in enumerate(
                    worksheet.iter_rows(min_row=1, max_row=20, values_only=True), 1,
                ):
                    keys = [
                        "".join(ch for ch in str(value or "").upper() if ch.isalnum())
                        for value in values
                    ]
                    code_col = next(
                        (i for i, key in enumerate(keys) if key in {"ITEMCODE", "CODE", "SOURCEITEMCODE"}),
                        None,
                    )
                    name_col = next(
                        (i for i, key in enumerate(keys) if key in {"ITEMNAME", "PRODUCTNAME", "DESCRIPTION"}),
                        None,
                    )
                    if code_col is not None and name_col is not None:
                        header_row = row_number
                        break
                if not header_row:
                    continue
                for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                    raw = str(values[code_col] or "").strip()
                    name = str(values[name_col] or "").strip()
                    if raw and name:
                        names.setdefault(norm_code(raw), name)
        except Exception:
            names = {}

    rows = []
    all_demand = list(run_row.get("uploaded_demand") or []) + list(
        run_row.get("fitting_demand") or []
    )
    for index, demand in enumerate(all_demand, 2):
        raw = str(demand.get("raw_code") or demand.get("item_code") or "")
        qty = float(demand.get("qty_pcs") or 0)
        rows.append(SimpleNamespace(
            source_row=index,
            raw_code=raw,
            item_name=names.get(norm_code(raw), ""),
            category=str(demand.get("material") or ""),
            qty_pcs=qty,
            file_weight_per_pc_kg=0.0,
            file_total_weight_kg=0.0,
        ))
    return rows


def _collect_blocks(
    engine_result,
    fitting_result,
    schedule_result,
    fitting_schedule,
    source_rows: Optional[Iterable],
) -> list[OperatorBlock]:
    names, source_raw = _source_maps(source_rows)
    item_map: dict[str, tuple[str, object]] = {}
    for plan_type, result in (("Pipe", engine_result), ("Fitting", fitting_result)):
        for item in (result.items if result else []):
            item_map[norm_code(item.item_code)] = (plan_type, item)

    unfinished: dict[str, object] = {}
    for schedule in (schedule_result, fitting_schedule):
        for item in (schedule.unfinished if schedule else []):
            unfinished[norm_code(item.item_code)] = item

    hours_left = {}
    for code, (_, item) in item_map.items():
        remaining = float(
            getattr(unfinished.get(code), "remaining_hours", 0.0) or 0.0
        )
        hours_left[code] = max(0.0, float(item.machine_hrs) - remaining)

    all_schedule_blocks = []
    for schedule in (schedule_result, fitting_schedule):
        if schedule:
            all_schedule_blocks.extend(schedule.blocks)
    max_day = max(
        (int(block.day) for block in all_schedule_blocks if not block.is_idle),
        default=0,
    )
    month = (
        getattr(engine_result, "effective_month", None)
        or getattr(fitting_result, "effective_month", None)
        or "1970-01"
    )
    dates = _plan_dates(max_day, month)
    output: list[OperatorBlock] = []

    for block in sorted(
        all_schedule_blocks,
        key=lambda value: (value.day, value.machine, value.shift, value.item_code or ""),
    ):
        if block.is_idle or not block.item_code:
            continue
        raw_hours = float(block.planned_hours)
        rounded_hours = int(round(raw_hours))
        if abs(raw_hours - rounded_hours) > 1e-7:
            raise OperatorExportInvariantError(
                f"{block.machine} day {block.day} has fractional block hours: {raw_hours}"
            )
        shift = str(block.shift or "").upper()
        if shift == "NIGHT" and rounded_hours != 10:
            raise OperatorExportInvariantError(
                f"{block.machine} day {block.day} NIGHT block is {rounded_hours}h, expected 10h"
            )
        code = norm_code(block.item_code)
        if code not in item_map:
            raise OperatorExportInvariantError(
                f"Schedule item {block.item_code} has no engine item"
            )
        plan_type, item = item_map[code]
        available = max(
            0.0, raw_hours - float(getattr(block, "excess_hours", 0.0) or 0.0)
        )
        productive = min(available, max(0.0, hours_left.get(code, 0.0)))
        hours_left[code] = max(0.0, hours_left.get(code, 0.0) - productive)
        if productive <= 0:
            continue
        fraction = productive / float(item.machine_hrs) if item.machine_hrs else 0.0
        raw_code = source_raw.get(code) or str(
            getattr(item, "raw_code", "") or item.item_code
        )
        output.append(OperatorBlock(
            machine=str(block.machine),
            plan_type=plan_type,
            day=int(block.day),
            date=dates[int(block.day)],
            shift=shift,
            block_hours=rounded_hours,
            item_code=str(item.item_code),
            raw_code=raw_code,
            item_name=names.get(code, ""),
            material=str(getattr(item, "material", "") or ""),
            target_pieces=float(item.qty_pcs) * fraction,
            kg=float(item.material_kg) * fraction,
            excess_hours=max(0.0, raw_hours - productive),
        ))
    return output


def _sheet_title(machine: str, used: set[str]) -> str:
    base = machine.replace("/", "-").replace("\\", "-")
    base = "".join(ch for ch in base if ch not in "[]:*?")[:31] or "Machine"
    title = base
    suffix = 2
    while title in used:
        marker = f"-{suffix}"
        title = f"{base[:31-len(marker)]}{marker}"
        suffix += 1
    used.add(title)
    return title


def _title(ws, text: str, subtitle: str, end_col: int) -> None:
    end = get_column_letter(end_col)
    ws.merge_cells(f"A1:{end}1")
    ws["A1"] = text
    ws["A1"].font = Font(name="Calibri", bold=True, size=15, color="FFFFFF")
    ws["A1"].fill = HEADER_FILL
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26
    ws.merge_cells(f"A2:{end}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Calibri", italic=True, size=9, color=NAVY)


def _headers(ws, row: int, values: list[str]) -> None:
    for column, value in enumerate(values, 1):
        cell = ws.cell(row, column, value)
        cell.fill = HEADER_FILL
        cell.font = Font(name="Calibri", bold=True, size=9, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def _print_setup(ws, header_row: int, landscape: bool = True) -> None:
    ws.freeze_panes = f"A{header_row + 1}"
    ws.print_title_rows = f"${header_row}:${header_row}"
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def operator_plan_bytes(
    engine_result=None,
    fitting_result=None,
    schedule_result=None,
    fitting_schedule=None,
    *,
    source_rows: Optional[Iterable] = None,
    unscheduled_rows: Optional[Iterable] = None,
) -> bytes:
    """Render the existing plan result in floor-operator workbook format."""
    month = (
        getattr(engine_result, "effective_month", None)
        or getattr(fitting_result, "effective_month", None)
        or "unknown"
    )
    blocks = _collect_blocks(
        engine_result, fitting_result, schedule_result, fitting_schedule, source_rows,
    )
    machines = sorted({block.machine for block in blocks})
    if not machines:
        raise OperatorExportInvariantError("Plan has no scheduled machine blocks")
    start_date = min(block.date for block in blocks)
    end_date = max(block.date for block in blocks)
    period = f"{start_date.day}–{end_date.day} {end_date.strftime('%B %Y')}"

    wb = Workbook()
    wb.remove(wb.active)

    # ① START HERE
    start = wb.create_sheet("① START HERE")
    _title(start, f"Machine Plan — {end_date.strftime('%B %Y')}", "How to use this workbook", 7)
    guidance = [
        ("What this is", f"A day-by-day production plan for each machine, for {period}."),
        ("Operators", "Go to your machine tab. It lists only your machine, by day, shift and block."),
        ("Supervisors", "Use ② Day Sheet for the morning walk-round."),
        ("Planners", "Use ③ Machine Loading for committed block hours and spare capacity."),
        ("Shifts", "DAY rows are orange; NIGHT rows are blue."),
        ("Block hours", "Run the whole committed block shown. Operator tabs never show fractional productive hours."),
        ("Night rule", "One item occupies the full 10-hour NIGHT block."),
        ("Minimum run", "Light blocks are marked “min-run block (light load)”."),
        ("Pieces", "Target Pieces is the existing scheduler result, rounded to whole pieces for display."),
    ]
    row = 4
    for label, text in guidance:
        start.cell(row, 1, label).font = Font(bold=True, color=NAVY)
        start.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        start.cell(row, 2, text).alignment = Alignment(wrap_text=True)
        row += 2
    row += 1
    _headers(start, row, ["Machine", "Type", "Days", "Items", "Block hours", "Pieces", "kg"])
    index_row = row + 1

    by_machine: dict[str, list[OperatorBlock]] = defaultdict(list)
    for block in blocks:
        by_machine[block.machine].append(block)
    used_titles = set(wb.sheetnames)
    machine_titles = {}
    for machine in machines:
        machine_titles[machine] = _sheet_title(machine, used_titles)

    for machine in machines:
        rows = by_machine[machine]
        ws = wb.create_sheet(machine_titles[machine])
        plan_types = ", ".join(sorted({block.plan_type for block in rows}))
        materials = ", ".join(sorted({block.material for block in rows if block.material}))
        _title(ws, f"MACHINE  {machine}", f"{plan_types} · {materials or 'Material not set'} · plan for {period}", 9)
        kpis = [
            ("Days scheduled", len({block.day for block in rows})),
            ("Different items", len({norm_code(block.item_code) for block in rows})),
            ("Machine hours", sum(block.block_hours for block in rows)),
            ("Total pieces", round(sum(block.target_pieces for block in rows))),
            ("Total kg", round(sum(block.kg for block in rows), 2)),
        ]
        for index, (label, value) in enumerate(kpis):
            column = index * 2 + 1
            ws.merge_cells(start_row=4, start_column=column, end_row=4, end_column=column + 1)
            ws.merge_cells(start_row=5, start_column=column, end_row=5, end_column=column + 1)
            ws.cell(4, column, label).font = Font(bold=True, size=8, color=NAVY)
            ws.cell(5, column, value).font = Font(bold=True, size=12, color=NAVY)
            ws.cell(5, column).alignment = Alignment(horizontal="center")
        headers = ["Day", "Date", "Shift", "Block\n(hours)", "Item Code", "Item Name", "Target Pieces", "kg", "Note"]
        _headers(ws, 7, headers)
        previous_day = None
        data_row = 8
        for block in rows:
            new_day = block.day != previous_day
            values = [
                block.day if new_day else None,
                block.date.strftime("%a %d %b %Y") if new_day else None,
                block.shift,
                block.block_hours,
                block.raw_code,
                block.item_name,
                round(block.target_pieces),
                round(block.kg, 2),
                "min-run block (light load)" if block.excess_hours > 0.5 else "",
            ]
            fill = NIGHT_FILL if block.shift == "NIGHT" else DAY_FILL
            for column, value in enumerate(values, 1):
                cell = ws.cell(data_row, column, value)
                cell.fill = fill
                cell.border = Border(
                    top=HEAVY if new_day else THIN, bottom=THIN, left=THIN, right=THIN
                )
                cell.alignment = Alignment(
                    horizontal="right" if column in (4, 7, 8) else "left",
                    vertical="center", wrap_text=column in (6, 9),
                )
                if column == 4:
                    cell.font = Font(bold=True, size=11, color=NAVY)
            previous_day = block.day
            data_row += 1
        total_row = data_row
        total_values = ["TOTAL", None, None, sum(b.block_hours for b in rows), None, None,
                        round(sum(b.target_pieces for b in rows)), round(sum(b.kg for b in rows), 2), None]
        for column, value in enumerate(total_values, 1):
            cell = ws.cell(total_row, column, value)
            cell.fill = TOTAL_FILL
            cell.font = Font(bold=True, color="FFFFFF")
            cell.border = Border(top=HEAVY, bottom=HEAVY)
        for column, width in enumerate([6, 17, 8, 11, 14, 34, 14, 12, 24], 1):
            ws.column_dimensions[get_column_letter(column)].width = width
        _print_setup(ws, 7)

        index_values = [
            machine, plan_types, len({b.day for b in rows}),
            len({norm_code(b.item_code) for b in rows}),
            sum(b.block_hours for b in rows),
            round(sum(b.target_pieces for b in rows)),
            round(sum(b.kg for b in rows), 2),
        ]
        for column, value in enumerate(index_values, 1):
            start.cell(index_row, column, value)
        start.cell(index_row, 1).hyperlink = f"#'{machine_titles[machine]}'!A1"
        start.cell(index_row, 1).style = "Hyperlink"
        index_row += 1
    for column, width in enumerate([22, 14, 10, 10, 14, 14, 14], 1):
        start.column_dimensions[get_column_letter(column)].width = width

    # ② Day Sheet
    day_sheet = wb.create_sheet("② Day Sheet")
    _title(day_sheet, "Day Sheet — what runs on each day",
           "All machines, grouped by date. Supervisor morning walk-round.", 8)
    _headers(day_sheet, 4, ["Date", "Shift", "Machine", "Block (hrs)", "Item Code", "Item Name", "Target Pieces", "kg"])
    row = 5
    previous_date = None
    for block in sorted(blocks, key=lambda b: (b.date, b.shift, b.machine, b.item_code)):
        if block.date != previous_date:
            day_sheet.cell(row, 1, block.date.strftime("%a %d %b %Y"))
            day_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            day_sheet.cell(row, 1).fill = PatternFill("solid", fgColor="DCE6F1")
            day_sheet.cell(row, 1).font = Font(bold=True, color=NAVY)
            row += 1
        values = [None, block.shift, block.machine, block.block_hours, block.raw_code,
                  block.item_name, round(block.target_pieces), round(block.kg, 2)]
        fill = NIGHT_FILL if block.shift == "NIGHT" else DAY_FILL
        for column, value in enumerate(values, 1):
            day_sheet.cell(row, column, value).fill = fill
        row += 1
        previous_date = block.date
    for column, width in enumerate([18, 9, 18, 12, 14, 36, 14, 12], 1):
        day_sheet.column_dimensions[get_column_letter(column)].width = width
    _print_setup(day_sheet, 4)

    # ③ Machine Loading
    loading = wb.create_sheet("③ Machine Loading")
    _title(loading, "Machine Loading — who is busy, who is free",
           "Committed block hours per day. Blank = idle.", 3 + len({b.date for b in blocks}))
    dates = sorted({block.date for block in blocks})
    headers = ["Machine", "Type"] + [date.strftime("%d-%b") for date in dates] + ["Total hrs"]
    _headers(loading, 4, headers)
    loading_totals = {}
    row = 5
    for machine in machines:
        machine_blocks = by_machine[machine]
        by_date = defaultdict(float)
        for block in machine_blocks:
            by_date[block.date] += block.block_hours
        loading_totals[machine] = sum(by_date.values())
        values = [machine, ", ".join(sorted({b.plan_type for b in machine_blocks}))] + [
            by_date.get(date) or None for date in dates
        ] + [sum(by_date.values())]
        for column, value in enumerate(values, 1):
            cell = loading.cell(row, column, value)
            if column > 2 and column <= 2 + len(dates) and value is not None:
                if value >= 16:
                    cell.fill = PatternFill("solid", fgColor="C6EFCE")
                elif value >= 8:
                    cell.fill = PatternFill("solid", fgColor="FFEB9C")
                else:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
        row += 1
    loading.column_dimensions["A"].width = 20
    loading.column_dimensions["B"].width = 12
    for column in range(3, len(headers) + 1):
        loading.column_dimensions[get_column_letter(column)].width = 11
    loading.freeze_panes = "C5"
    loading.print_title_rows = "$4:$4"
    loading.page_setup.orientation = "landscape"
    loading.page_setup.fitToWidth = 1

    # ④ Item Plan
    item_sheet = wb.create_sheet("④ Item Plan")
    _title(item_sheet, "Item Plan — where each item is made",
           "One row per item: machines, days, total pieces and kg.", 7)
    _headers(item_sheet, 4, ["Item Code", "Item Name", "Material", "Machine(s)", "Days", "Total Pieces", "Total kg"])
    by_item = defaultdict(list)
    for block in blocks:
        by_item[norm_code(block.item_code)].append(block)
    row = 5
    for item_blocks in sorted(by_item.values(), key=lambda values: values[0].raw_code):
        first = item_blocks[0]
        values = [
            first.raw_code, first.item_name, first.material,
            ", ".join(sorted({b.machine for b in item_blocks})),
            len({b.day for b in item_blocks}),
            round(sum(b.target_pieces for b in item_blocks)),
            round(sum(b.kg for b in item_blocks), 2),
        ]
        for column, value in enumerate(values, 1):
            item_sheet.cell(row, column, value)
        row += 1
    item_sheet.auto_filter.ref = f"A4:G{max(4, row - 1)}"
    for column, width in enumerate([14, 36, 12, 30, 9, 14, 12], 1):
        item_sheet.column_dimensions[get_column_letter(column)].width = width
    _print_setup(item_sheet, 4)

    # ⑤ Not Scheduled
    not_scheduled = wb.create_sheet("⑤ Not Scheduled")
    _headers(not_scheduled, 1, ["Source Row", "Item Code", "Item Name", "Category", "Required Pieces", "Weight kg", "Reason"])
    row = 2
    for entry in unscheduled_rows or []:
        stock = getattr(entry, "stock", entry)
        values = [
            getattr(stock, "source_row", ""),
            getattr(stock, "raw_code", ""),
            getattr(stock, "item_name", ""),
            getattr(stock, "category", ""),
            getattr(stock, "qty_pcs", 0),
            getattr(stock, "file_total_weight_kg", 0),
            getattr(entry, "reason", ""),
        ]
        for column, value in enumerate(values, 1):
            not_scheduled.cell(row, column, value)
        row += 1
    not_scheduled.auto_filter.ref = f"A1:G{max(1, row - 1)}"
    not_scheduled.freeze_panes = "A2"

    # ⑥ Plan Summary
    summary = wb.create_sheet("⑥ Plan Summary")
    metrics = [
        ("Plan month", month),
        ("Machines", len(machines)),
        ("Scheduled blocks", len(blocks)),
        ("Committed block hours", sum(block.block_hours for block in blocks)),
        ("Scheduled pieces", round(sum(block.target_pieces for block in blocks))),
        ("Scheduled kg", round(sum(block.kg for block in blocks), 2)),
        ("Completion date", end_date),
        ("Persistence", "Export only; workbook generation does not modify the plan"),
    ]
    _headers(summary, 1, ["Metric", "Value"])
    for row, (label, value) in enumerate(metrics, 2):
        summary.cell(row, 1, label)
        summary.cell(row, 2, value)
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 42
    summary.freeze_panes = "A2"

    # ⑦ Source Stock / demand
    source = wb.create_sheet("⑦ Source Stock")
    _headers(source, 1, ["Source Row", "Item Code", "Match Key", "Item Name", "Category", "Required Pieces", "Weight/Pc kg", "Total Weight kg"])
    row = 2
    if source_rows:
        for item in source_rows:
            raw = str(getattr(item, "raw_code", "") or "")
            values = [
                getattr(item, "source_row", ""),
                raw, norm_code(raw), getattr(item, "item_name", ""),
                getattr(item, "category", ""), getattr(item, "qty_pcs", 0),
                getattr(item, "file_weight_per_pc_kg", 0),
                getattr(item, "file_total_weight_kg", 0),
            ]
            for column, value in enumerate(values, 1):
                source.cell(row, column, value)
            row += 1
    else:
        export_items = []
        for result in (engine_result, fitting_result):
            for item in (result.items if result else []):
                export_items.append(item)
        for item in sorted(export_items, key=lambda value: value.item_code):
            raw = str(getattr(item, "raw_code", "") or item.item_code)
            values = ["", raw, norm_code(raw), "", getattr(item, "material", ""),
                      getattr(item, "qty_pcs", 0), getattr(item, "weight_per_pc_kg", 0),
                      getattr(item, "material_kg", 0)]
            for column, value in enumerate(values, 1):
                source.cell(row, column, value)
            row += 1
    source.auto_filter.ref = f"A1:H{max(1, row - 1)}"
    source.freeze_panes = "A2"

    # Final invariants.
    actual_machine_tabs = [
        title for title in wb.sheetnames
        if title not in {"① START HERE", "② Day Sheet", "③ Machine Loading",
                         "④ Item Plan", "⑤ Not Scheduled", "⑥ Plan Summary", "⑦ Source Stock"}
    ]
    if len(actual_machine_tabs) != len(machines):
        raise OperatorExportInvariantError("Machine worksheet count does not match plan")
    for machine, total in loading_totals.items():
        if total != sum(block.block_hours for block in by_machine[machine]):
            raise OperatorExportInvariantError(f"Machine Loading does not reconcile for {machine}")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()