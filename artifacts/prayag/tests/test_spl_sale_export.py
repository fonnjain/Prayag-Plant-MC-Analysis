import types

from openpyxl import load_workbook
import pytest

from spl_sale_export import (
    ClassifiedRow,
    SplSalePlan,
    StockCodeAmbiguityError,
    StockRow,
    classify_stock,
    export_plan,
)


def _stock(code: str, qty: float = 1.0, category: str = "CPVC-PIPE") -> StockRow:
    return StockRow(2, code, code, category, 0.0, qty, 1.0, qty)


def _route(code: str, machine: str = "M/C-1") -> dict:
    return {
        "item_code": code,
        "machine": machine,
        "material": "CPVC",
        "capable": True,
    }


def test_classification_matches_through_shared_normaliser():
    scheduled, unscheduled = classify_stock(
        [_stock("PS-2"), _stock("PS-2S")],
        [_route("PS2"), _route("PS2S")],
    )
    assert not unscheduled
    assert [row.routing_code for row in scheduled] == ["PS2", "PS2S"]
    assert scheduled[0].stock.match_code != scheduled[1].stock.match_code


def test_ambiguous_normalised_routes_are_not_guessed():
    scheduled, unscheduled = classify_stock(
        [_stock("PS-2")],
        [_route("PS2"), _route("PS-2", "M/C-2")],
    )
    assert not scheduled
    assert len(unscheduled) == 1
    assert unscheduled[0].reason.startswith("Ambiguous")


def test_genuine_gap_and_blank_subtotal_stay_unscheduled():
    scheduled, unscheduled = classify_stock(
        [_stock("NO-ROUTE"), _stock("Cpvc Fitting", 0, "")],
        [_route("PS2")],
    )
    assert not scheduled
    assert [row.reason for row in unscheduled] == [
        "No capable Plumbing machine route",
        "Blank-category subtotal/non-item row; column E is zero",
    ]


def test_duplicate_normalised_stock_codes_fail_loudly():
    with pytest.raises(StockCodeAmbiguityError, match="PS2"):
        classify_stock(
            [_stock("PS-2"), _stock("PS.2")],
            [_route("PS2")],
        )


def _coverage(no_weight=None, no_machine=None):
    return types.SimpleNamespace(
        no_weight=no_weight or [], no_machine=no_machine or [],
    )


def _schedule(blocks=None, unfinished=None):
    return types.SimpleNamespace(
        week_days=[1], blocks=blocks or [], unfinished=unfinished or [],
    )


def test_missing_bom_item_exports_without_blocks_or_false_completion(tmp_path):
    stock = _stock("A-721", 190, "AGRI-FG")
    routed = ClassifiedRow(
        stock, "scheduled", "", "A721", (_route("A721", "A03(U-150)"),),
    )
    item = types.SimpleNamespace(
        item_code="A721", raw_code="A-721", material="AGRI", qty_pcs=190.0,
        material_kg=0.0, machine_hrs=0.0,
    )
    empty_engine = types.SimpleNamespace(
        effective_month="2026-08", items=[], coverage_gaps=_coverage(),
    )
    plan = SplSalePlan(
        [stock], [routed], [], [], [], empty_engine,
        types.SimpleNamespace(
            effective_month="2026-08", items=[item],
            coverage_gaps=_coverage(no_weight=["A721"]),
        ),
        _schedule(), _schedule(),
    )
    output = export_plan(plan, tmp_path / "a721.xlsx")
    workbook = load_workbook(output, read_only=True, data_only=True)
    summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
    item_row = next(workbook["Item Summary"].iter_rows(min_row=2, values_only=True))
    assert summary["Fully scheduled items"] == 0
    assert summary["Completion date"] is None
    assert workbook["Daily Machine Plan"].max_row == 1
    assert item_row[-1] == "Routed; not scheduled — missing BOM weight or machine input"


def test_unfinished_item_is_reported_as_partial(tmp_path):
    stock = _stock("PS-2", 100)
    routed = ClassifiedRow(stock, "scheduled", "", "PS2", (_route("PS2"),))
    item = types.SimpleNamespace(
        item_code="PS2", raw_code="PS-2", material="CPVC", qty_pcs=100.0,
        material_kg=10.0, machine_hrs=10.0,
    )
    block = types.SimpleNamespace(
        is_idle=False, item_code="PS2", planned_hours=5.0, excess_hours=0.0,
        day=1, week=1, machine="M/C-1", shift="DAY",
    )
    remaining = types.SimpleNamespace(
        item_code="PS2", remaining_pcs=50.0, remaining_hours=5.0,
    )
    pipe_engine = types.SimpleNamespace(
        effective_month="2026-08", items=[item], coverage_gaps=_coverage(),
    )
    empty_engine = types.SimpleNamespace(
        effective_month="2026-08", items=[], coverage_gaps=_coverage(),
    )
    plan = SplSalePlan(
        [stock], [routed], [], [], [], pipe_engine, empty_engine,
        _schedule([block], [remaining]), _schedule(),
    )
    output = export_plan(plan, tmp_path / "partial.xlsx")
    workbook = load_workbook(output, read_only=True, data_only=True)
    summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
    item_row = next(workbook["Item Summary"].iter_rows(min_row=2, values_only=True))
    assert summary["Fully scheduled items"] == 0
    assert summary["Partially scheduled items"] == 1
    assert summary["Physically scheduled pieces"] == 50
    assert item_row[-1].startswith("Partially scheduled — 50.000 of 100.000 pcs")


def test_complete_item_with_missing_block_hours_fails_export(tmp_path):
    stock = _stock("PS-2", 100)
    routed = ClassifiedRow(stock, "scheduled", "", "PS2", (_route("PS2"),))
    item = types.SimpleNamespace(
        item_code="PS2", raw_code="PS-2", material="CPVC", qty_pcs=100.0,
        material_kg=10.0, machine_hrs=10.0,
    )
    short_block = types.SimpleNamespace(
        is_idle=False, item_code="PS2", planned_hours=1.0, excess_hours=0.0,
        day=1, week=1, machine="M/C-1", shift="DAY",
    )
    pipe_engine = types.SimpleNamespace(
        effective_month="2026-08", items=[item], coverage_gaps=_coverage(),
    )
    empty_engine = types.SimpleNamespace(
        effective_month="2026-08", items=[], coverage_gaps=_coverage(),
    )
    plan = SplSalePlan(
        [stock], [routed], [], [], [], pipe_engine, empty_engine,
        _schedule([short_block]), _schedule(),
    )
    with pytest.raises(ValueError, match="do not cover required hours"):
        export_plan(plan, tmp_path / "short.xlsx")


def test_partial_item_with_missing_scheduled_hours_fails_export(tmp_path):
    stock = _stock("PS-2", 100)
    routed = ClassifiedRow(stock, "scheduled", "", "PS2", (_route("PS2"),))
    item = types.SimpleNamespace(
        item_code="PS2", raw_code="PS-2", material="CPVC", qty_pcs=100.0,
        material_kg=10.0, machine_hrs=10.0,
    )
    short_block = types.SimpleNamespace(
        is_idle=False, item_code="PS2", planned_hours=1.0, excess_hours=0.0,
        day=1, week=1, machine="M/C-1", shift="DAY",
    )
    remaining = types.SimpleNamespace(
        item_code="PS2", remaining_pcs=50.0, remaining_hours=5.0,
    )
    pipe_engine = types.SimpleNamespace(
        effective_month="2026-08", items=[item], coverage_gaps=_coverage(),
    )
    empty_engine = types.SimpleNamespace(
        effective_month="2026-08", items=[], coverage_gaps=_coverage(),
    )
    plan = SplSalePlan(
        [stock], [routed], [], [], [], pipe_engine, empty_engine,
        _schedule([short_block], [remaining]), _schedule(),
    )
    with pytest.raises(ValueError, match="do not cover required hours"):
        export_plan(plan, tmp_path / "partial-short.xlsx")